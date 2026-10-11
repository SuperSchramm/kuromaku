#!/usr/bin/env python3
"""
phase-report.py — Kuromaku Phase 5: Report Generation

Two report formats:
  1. (default) CCSO-style full assessment report, rendered from
     assets/report-template.docx via docxtpl — a professional pentest-
     report layout, with Risk Score/Exploitation Likelihood/Business
     Impact per finding DERIVED FROM nuclei's own CVSS data (vector +
     score), terminology aligned with Bugcrowd's Vulnerability Rating
     Taxonomy (https://bugcrowd.com/vulnerability-rating-taxonomy) —
     not a fabricated judgment call. Remediation Difficulty is
     deliberately NOT included: unlike CVSS/VRT-derivable fields, it
     depends on the target's own infrastructure (in-house skills,
     hardware budget, change-control process) that an external scan has
     no way to know.
  2. Legacy flat summary report (the original kuromaku format) — used as
     a fallback if docxtpl / the template asset / python-docx aren't
     available, or if the new path throws for ANY reason. The report
     phase should never fail outright just because the richer format
     couldn't be built — same resilience philosophy as the LLM-summary
     fallback below.

Usage:
  phase-report.py -d <domain> -t <scan_dir> [--llm-url URL] [--llm-model NAME]
                   [--team-name NAME] [--client-name NAME]

If --llm-url is unreachable, the executive summary falls back to a
templated paragraph — the report still generates fully either way.

Required inputs (whatever exists is used; missing files are skipped gracefully):
  <prefix>-uniqdomains.log
  <prefix>-naabu.log
  <prefix>-resolved.log
  <prefix>-nucleiAlerts.log     (JSON LINES — phase-webscan.sh's nuclei runs
                                  with -jsonl; this is NOT the old plain-text
                                  "[template] [protocol] [severity] target"
                                  format kuromaku used before)
  <prefix>-dirsearch.log
  <prefix>-dalfox.log
  <prefix>-nmapvulners.log
  <prefix>-out-of-scope.log
  <tdir>/scope.txt               (written by the orchestrator when `scope`
                                   was passed to kuromaku_run — feeds the
                                   templated report's Scope > Networks table)

Output:
  <prefix>-report.docx
"""

import argparse
import json
import os
import re
import sys
from collections import defaultdict
from datetime import datetime

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    HAVE_DOCX = True
except ImportError:
    HAVE_DOCX = False

try:
    from docxtpl import DocxTemplate, InlineImage
    from docx.shared import Mm
    HAVE_DOCXTPL = True
except ImportError:
    HAVE_DOCXTPL = False

TEMPLATE_PATH = os.environ.get(
    "REPORT_TEMPLATE_PATH",
    os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets", "report-template.docx"),
)

# ─── Severity reference tables ──────────────────────────────────────────────
SEVERITY_ORDER = ["critical", "high", "medium", "low", "info"]
SEVERITY_COLORS = {
    "critical": RGBColor(0xC0, 0x00, 0x00) if HAVE_DOCX else None,
    "high":     RGBColor(0xE3, 0x6C, 0x09) if HAVE_DOCX else None,
    "medium":   RGBColor(0xBF, 0x8F, 0x00) if HAVE_DOCX else None,
    "low":      RGBColor(0x38, 0x76, 0xC0) if HAVE_DOCX else None,
    "info":     RGBColor(0x70, 0x70, 0x70) if HAVE_DOCX else None,
}
# Matches the template's OWN "Risk Classifications" legend (Critical=10,
# High=7-9, Medium=4-6, Low=1-3, Informational=0) — used as the Risk Score
# fallback for findings nuclei didn't attach a CVSS score to (common for
# non-CVE templates: exposures/misconfiguration/technologies/etc).
SEVERITY_RISK_SCORE_FALLBACK = {"critical": 10, "high": 8, "medium": 5, "low": 2, "info": 0}
SEVERITY_RISK_LABEL = {
    "critical": "Critical", "high": "High", "medium": "Medium",
    "low": "Low", "info": "Informational",
}
# Fallback Likelihood/Impact when a finding has no CVSS v3 vector to derive
# them from — a coarser, severity-based approximation, not a guess.
SEVERITY_LIKELIHOOD_FALLBACK = {
    "critical": "Likely", "high": "Possible", "medium": "Possible",
    "low": "Unlikely", "info": "Unlikely",
}
SEVERITY_IMPACT_FALLBACK = {
    "critical": "Major", "high": "Major", "medium": "Moderate",
    "low": "Minor", "info": "Minor",
}


def parse_args():
    p = argparse.ArgumentParser()
    p.add_argument("-d", "--domain")
    p.add_argument("-dl", "--domain_list")
    p.add_argument("-t", "--target-dir", required=True)
    p.add_argument("--llm-url", default=os.environ.get("LLM_URL", "http://host.docker.internal:1234/v1"))
    p.add_argument("--llm-model", default=os.environ.get("LLM_MODEL", "qwen"))
    p.add_argument("--team-name", default=os.environ.get("REPORT_TEAM_NAME", "Kuromaku Security Assessment"))
    p.add_argument("--client-name", default=os.environ.get("REPORT_CLIENT_NAME", "<CLIENT NAME>"))
    p.add_argument("--legacy", action="store_true", help="Force the old flat-summary format, skip the templated report entirely.")
    return p.parse_args()


def read_lines(path):
    if not os.path.exists(path):
        return []
    with open(path, errors="ignore") as f:
        return [l.rstrip("\n") for l in f if l.strip()]


# ─── nuclei JSONL parsing ────────────────────────────────────────────────────
# phase-webscan.sh runs nuclei with -jsonl (see CLAUDE.md / that script) —
# nucleiAlerts.log is one JSON object per line, NOT the old plain-text
# "[template] [protocol] [severity] target" format. Defensive: skip any line
# that isn't valid JSON rather than crash the whole report over one bad line
# (nuclei's own banner/stats noise should never reach -o with -silent set,
# but malformed/partial lines from a killed/timed-out batch are plausible).
def parse_nuclei_jsonl(lines):
    findings = []
    for line in lines:
        line = line.strip()
        if not line or not line.startswith("{"):
            continue
        try:
            rec = json.loads(line)
        except json.JSONDecodeError:
            continue
        info = rec.get("info") or {}
        classification = info.get("classification") or {}
        severity = (info.get("severity") or "info").lower()
        if severity not in SEVERITY_RISK_LABEL:
            severity = "info"
        references = info.get("reference") or []
        if isinstance(references, str):
            references = [references]
        findings.append({
            "template_id": rec.get("template-id", "unknown"),
            "name": info.get("name") or rec.get("template-id") or "Unknown Finding",
            "severity": severity,
            "cvss_score": classification.get("cvss-score"),
            "cvss_vector": classification.get("cvss-metrics"),
            "description": (info.get("description") or "").strip(),
            "remediation": (info.get("remediation") or "").strip(),
            "references": references,
            "host": rec.get("host") or rec.get("ip") or "",
            "matched_at": rec.get("matched-at") or rec.get("host") or "",
            "extracted_results": rec.get("extracted-results") or [],
        })
    return findings


# ─── CVSS v3 vector -> Likelihood / Business Impact ─────────────────────────
# Grounded in nuclei's own CVSS data (the same data Bugcrowd's VRT and CVSS
# itself are built on), not an LLM guess or an invented heuristic. Only
# handles CVSS v3.x vectors (the format nuclei's current templates use) —
# an older v2 vector, or no vector at all (common for non-CVE templates),
# falls back to the severity-based tables above in build_finding_context().
def parse_cvss_vector(vector):
    if not vector or "CVSS:3" not in vector:
        return {}
    metrics = {}
    for part in vector.split("/"):
        if ":" not in part:
            continue
        k, _, v = part.partition(":")
        metrics[k] = v
    return metrics


def derive_likelihood(metrics):
    if not metrics:
        return None
    av, ac, pr, ui = metrics.get("AV"), metrics.get("AC"), metrics.get("PR"), metrics.get("UI")
    if av == "N" and ac == "L" and pr == "N" and ui == "N":
        return "Likely"
    if ac == "H" or pr == "H" or av in ("L", "P"):
        return "Unlikely"
    return "Possible"


def derive_business_impact(metrics):
    if not metrics:
        return None
    vals = [metrics.get(k) for k in ("C", "I", "A") if metrics.get(k)]
    if not vals:
        return None
    if "H" in vals:
        return "Major"
    if all(v == "N" for v in vals):
        return "Minor"
    return "Moderate"


def build_finding_context(findings):
    """Sort by CVSS score (falling back to the severity-based score table)
    descending, and shape each finding for the template. Every auto-filled
    field here traces to real nuclei data — CVSS vector components, or the
    template's own documented description/remediation/reference fields —
    nothing is fabricated analyst judgment."""
    def sort_key(f):
        return f["cvss_score"] if f["cvss_score"] is not None else SEVERITY_RISK_SCORE_FALLBACK.get(f["severity"], 0)

    ordered = sorted(findings, key=sort_key, reverse=True)
    out = []
    for f in ordered:
        metrics = parse_cvss_vector(f["cvss_vector"])
        risk_score = f["cvss_score"] if f["cvss_score"] is not None else SEVERITY_RISK_SCORE_FALLBACK.get(f["severity"], 0)
        likelihood = derive_likelihood(metrics) or SEVERITY_LIKELIHOOD_FALLBACK.get(f["severity"], "Possible")
        business_impact = derive_business_impact(metrics) or SEVERITY_IMPACT_FALLBACK.get(f["severity"], "Moderate")

        security_implications = f["description"] or (
            f"Automated detection via nuclei template '{f['template_id']}' "
            f"({SEVERITY_RISK_LABEL.get(f['severity'], f['severity'])} severity). "
            "Manual verification is recommended to confirm impact."
        )

        analysis_parts = []
        if f["matched_at"]:
            analysis_parts.append(f"Matched at: {f['matched_at']}")
        if f["extracted_results"]:
            analysis_parts.append("Extracted: " + "; ".join(str(x) for x in f["extracted_results"][:5]))
        if f["cvss_vector"]:
            analysis_parts.append(f"CVSS vector: {f['cvss_vector']}")
        analysis = " ".join(analysis_parts) or "See raw nuclei output (nucleiAlerts.log) for match details."

        recommendation = f["remediation"] or (
            f"Review this finding against the template documentation at "
            f"https://github.com/projectdiscovery/nuclei-templates (template: "
            f"{f['template_id']}) and remediate per vendor guidance for the "
            f"affected technology."
        )

        references_text = " | ".join(f["references"][:5]) if f["references"] else (
            "https://github.com/projectdiscovery/nuclei-templates"
        )

        rs = round(risk_score, 1) if isinstance(risk_score, float) else risk_score
        out.append({
            "title": f["name"],
            "risk_score": rs,
            "risk_label": SEVERITY_RISK_LABEL.get(f["severity"], f["severity"].title()),
            "likelihood": likelihood,
            "business_impact": business_impact,
            "security_implications": security_implications,
            "analysis": analysis,
            "recommendation": recommendation,
            "references_text": references_text,
        })
    return out


# ─── naabu / dirsearch parsing (shared by both report formats) ─────────────
def parse_naabu(lines):
    results = []
    for line in lines:
        if ":" in line:
            host, _, port = line.rpartition(":")
            results.append((host.strip(), port.strip()))
    return results


DIRSEARCH_LINE_RE = re.compile(r"^\s*(\d{3})\s+\S+\s+(\S+)")

def parse_dirsearch(lines, interesting_codes=("200", "301", "302", "401", "403")):
    results = []
    for line in lines:
        m = DIRSEARCH_LINE_RE.match(line)
        if m and m.group(1) in interesting_codes:
            results.append((m.group(1), m.group(2)))
    return results


# ─── LLM call (optional, shared) ────────────────────────────────────────────
def generate_executive_summary(llm_url, llm_model, target, sev_counts, top_findings):
    """Call LM Studio's OpenAI-compatible API for a short exec summary.
    Falls back to a templated paragraph on any failure — including
    host.docker.internal not resolving; see orchestrator.js's --add-host."""
    fallback = (
        f"This report presents the findings of an automated reconnaissance and "
        f"vulnerability assessment conducted against {target}. The assessment "
        f"identified {sum(sev_counts.values())} findings across "
        f"{len([k for k in SEVERITY_ORDER if sev_counts.get(k)])} severity levels, "
        f"including {sev_counts.get('critical', 0)} critical and "
        f"{sev_counts.get('high', 0)} high severity issues. "
        f"Findings should be triaged in order of severity, with critical and high "
        f"items addressed first given their potential for unauthenticated exploitation."
    )

    try:
        import urllib.request

        prompt = (
            f"Write a concise 3-4 sentence executive summary paragraph for a "
            f"security assessment report on the target '{target}'. "
            f"The scan found: {sev_counts.get('critical',0)} critical, "
            f"{sev_counts.get('high',0)} high, {sev_counts.get('medium',0)} medium, "
            f"{sev_counts.get('low',0)} low, and {sev_counts.get('info',0)} "
            f"informational findings. "
            f"Top issues include: {', '.join(top_findings[:3]) if top_findings else 'none notable'}. "
            f"Write in formal security-report tone. Do not invent specific CVE numbers "
            f"or details beyond what is given. Output only the paragraph, no headers."
        )

        body = json.dumps({
            "model": llm_model,
            "max_tokens": 300,
            "messages": [{"role": "user", "content": prompt}],
            "temperature": 0.3,
        }).encode()

        req = urllib.request.Request(
            f"{llm_url}/chat/completions",
            data=body,
            headers={"Content-Type": "application/json"},
        )
        with urllib.request.urlopen(req, timeout=30) as resp:
            data = json.loads(resp.read())
            text = data["choices"][0]["message"]["content"].strip()
            if text:
                return text
    except Exception as e:
        print(f"[report] LLM summary generation failed ({e}) — using template", file=sys.stderr)

    return fallback


# ─── Methodology diagram (replaces the template's generic manual-pentest
# hexagon with kuromaku's actual pipeline). Optional — any failure (Pillow
# missing, font missing, draw error) just means the report renders without
# this image rather than crashing. ──────────────────────────────────────────
def generate_methodology_diagram(out_path):
    try:
        from PIL import Image, ImageDraw, ImageFont
    except ImportError:
        print("[report] Pillow not installed — skipping methodology diagram", file=sys.stderr)
        return False

    try:
        W, H = 1400, 420
        img = Image.new("RGB", (W, H), "white")
        draw = ImageDraw.Draw(img)

        font_title = font_body = font_note = None
        for candidate_dir in ("/usr/share/fonts/truetype/dejavu", "/usr/share/fonts/dejavu"):
            try:
                font_title = ImageFont.truetype(f"{candidate_dir}/DejaVuSans-Bold.ttf", 26)
                font_body = ImageFont.truetype(f"{candidate_dir}/DejaVuSans.ttf", 16)
                font_note = ImageFont.truetype(f"{candidate_dir}/DejaVuSans-Oblique.ttf", 15)
                break
            except Exception:
                continue
        if font_title is None:
            font_title = font_body = font_note = ImageFont.load_default()

        phases = [
            ("Recon", ["subfinder / amass", "gau / katana"]),
            ("Network", ["dnsx resolution", "naabu port scan"]),
            ("Webscan", ["nuclei (14 categories)", "nmap+vulners, dirsearch"]),
            ("XSS", ["dalfox on reflected", "parameters"]),
            ("Report", ["this DOCX", "assessment report"]),
        ]

        box_w, box_h = 220, 150
        gap = 55
        total_w = len(phases) * box_w + (len(phases) - 1) * gap
        start_x = (W - total_w) // 2
        y = 50

        fill_color = (0x38, 0x76, 0xC0)
        arrow_color = (0x50, 0x50, 0x50)

        def text_w(s, font):
            try:
                return draw.textlength(s, font=font)
            except AttributeError:
                return draw.textsize(s, font=font)[0]

        for i, (title, lines) in enumerate(phases):
            x = start_x + i * (box_w + gap)
            try:
                draw.rounded_rectangle([x, y, x + box_w, y + box_h], radius=14, fill=fill_color)
            except AttributeError:
                draw.rectangle([x, y, x + box_w, y + box_h], fill=fill_color)
            tw = text_w(title, font_title)
            draw.text((x + (box_w - tw) / 2, y + 14), title, fill="white", font=font_title)
            for li, line in enumerate(lines):
                lw = text_w(line, font_body)
                draw.text((x + (box_w - lw) / 2, y + 56 + li * 24), line, fill="white", font=font_body)
            if i < len(phases) - 1:
                ax1, ax2 = x + box_w, x + box_w + gap
                ay = y + box_h // 2
                draw.line([(ax1 + 6, ay), (ax2 - 8, ay)], fill=arrow_color, width=3)
                draw.polygon([(ax2 - 8, ay - 8), (ax2 + 2, ay), (ax2 - 8, ay + 8)], fill=arrow_color)

        note = "IP/CIDR-mode scans skip Recon and XSS (no hostnames) and run only Network → Webscan."
        nw = text_w(note, font_note)
        draw.text(((W - nw) / 2, y + box_h + 25), note, fill=(0x60, 0x60, 0x60), font=font_note)

        img.save(out_path, "PNG")
        return True
    except Exception as e:
        print(f"[report] methodology diagram generation failed ({e}) — skipping", file=sys.stderr)
        return False


# ─── Templated (docxtpl) report ─────────────────────────────────────────────
# ─── Assessment-data appendix for the templated report ───────────────────────
# The CCSO template only has sections for scope, methodology and nuclei
# findings. Everything else the pipeline collects — hosts, IPs, open ports,
# nmap service/CVE results, dirsearch, dalfox — was read and counted but never
# rendered, so reports looked like an empty template. These sections are
# appended after the template renders, using only styles the template has
# (Normal, Heading 1-3) and explicit table borders.
ANSI_RE = re.compile(r"\x1b\[[0-9;]*[A-Za-z]|\[\d+(?:;\d+)*m")


def strip_ansi(text):
    return ANSI_RE.sub("", text)


def parse_nmap(lines):
    """Per-host results from nmap's normal (-oN) output."""
    hosts = []
    cur = None
    for raw in lines:
        line = raw.rstrip()
        m = re.match(r"^Nmap scan report for (\S+)(?: \(([\d.:a-fA-F]+)\))?", line)
        if m:
            cur = {"host": m.group(1), "ip": m.group(2) or m.group(1),
                   "timeout": False, "ports": [], "cves": []}
            hosts.append(cur)
            continue
        m = re.match(r"^Skipping host (\S+) .*due to host timeout", line)
        if m:
            for h in hosts:
                if h["host"] == m.group(1):
                    h["timeout"] = True
            continue
        if cur is None:
            continue
        m = re.match(r"^(\d+/(?:tcp|udp))\s+(\S+)\s+(\S+)\s*(.*)$", line)
        if m:
            cur["ports"].append((m.group(1), m.group(3), m.group(4).strip()))
            continue
        for cve in re.findall(r"CVE-\d{4}-\d{4,7}", line):
            if cve not in cur["cves"]:
                cur["cves"].append(cve)
    return hosts


def _set_table_borders(table):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "808080")
        borders.append(el)
    # CT_TblPr is an ordered sequence: tblBorders must come BEFORE shd,
    # tblLayout, tblCellMar and tblLook. Appending it after tblLook (what
    # python-docx leaves at the end) makes Word report unreadable content.
    successor = None
    for tag in ("shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription"):
        successor = tbl_pr.find(qn(f"w:{tag}"))
        if successor is not None:
            break
    if successor is not None:
        successor.addprevious(borders)
    else:
        tbl_pr.append(borders)


def _add_table(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    _set_table_borders(table)
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        cell.paragraphs[0].add_run(h).bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    return table


def _add_grid(document, items, cols):
    """Short strings laid out row-wise in `cols` columns (hostnames, IPs)."""
    rows = [items[i:i + cols] for i in range(0, len(items), cols)]
    table = document.add_table(rows=0, cols=cols)
    _set_table_borders(table)
    for r in rows:
        cells = table.add_row().cells
        for i in range(cols):
            cells[i].text = r[i] if i < len(r) else ""
    return table


def dedupe_paragraph_ids(out_path):
    """Make every w14:paraId in the body unique.

    docxtpl repeats paragraph/table blocks for each finding; the copies carry
    the original w14:paraId, so a report with N findings holds N paragraphs
    with the same id (the source template already has some). Word treats a
    duplicate paraId as unreadable content and offers to recover the file.
    """
    from docx import Document as _Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    document = _Document(out_path)

    # Styles without a <w:name>: the CCSO template (a Google Docs export)
    # carries 14 table styles ("a".."ac") with no name element. Word cannot
    # register a nameless style, so it opens the file with a "Show Repairs:
    # Styles" recovery prompt. Give each a unique name; the styleId that
    # table properties reference is untouched.
    named = 0
    for style in document.styles.element.findall(qn("w:style")):
        if style.find(qn("w:name")) is None:
            name_el = OxmlElement("w:name")
            name_el.set(qn("w:val"), f"Kuromaku Table {style.get(qn('w:styleId'))}")
            style.insert(0, name_el)           # schema order: name comes first
            named += 1

    attr = "{http://schemas.microsoft.com/office/word/2010/wordml}paraId"
    seen = set()
    changed = 0
    counter = 0x10000000
    for el in document.element.body.iter():
        pid = el.get(attr)
        if pid is None:
            continue
        if pid in seen:
            counter += 1
            pid = f"{counter:08X}"          # < 0x80000000 as the spec requires
            el.set(attr, pid)
            changed += 1
        seen.add(pid)
    if changed or named:
        document.save(out_path)
    return changed + named


def append_assessment_data(out_path, uniqdomains, unique_ips, naabu_results,
                           nmap_hosts, nmap_raw_lines, dirsearch_raw, dirsearch_results,
                           dalfox_raw, dalfox_findings):
    from docx import Document as _Document
    document = _Document(out_path)
    cap = 300

    document.add_page_break()
    document.add_heading("APPENDIX C - SCAN COVERAGE AND RAW RESULTS", level=1)
    document.add_paragraph(
        "This appendix lists what the automated pipeline discovered and tested, "
        "including areas where nothing was found, so that coverage can be reviewed "
        "alongside the findings above."
    )

    document.add_heading("C.1 Hostnames in Scope", level=2)
    if uniqdomains:
        document.add_paragraph(f"{len(uniqdomains)} unique in-scope hostname(s) were discovered.")
        _add_grid(document, uniqdomains[:cap], 3)
        if len(uniqdomains) > cap:
            document.add_paragraph(f"... and {len(uniqdomains) - cap} more (see uniqdomains.log)")
    else:
        document.add_paragraph("No hostnames were recorded for this target.")

    document.add_heading("C.2 Resolved IP Addresses", level=2)
    if unique_ips:
        document.add_paragraph(f"{len(unique_ips)} unique IP address(es) resolved.")
        _add_grid(document, unique_ips[:cap], 4)
        if len(unique_ips) > cap:
            document.add_paragraph(f"... and {len(unique_ips) - cap} more (see unique-ips.log)")
    else:
        document.add_paragraph("No IP addresses were recorded for this target.")

    document.add_heading("C.3 Open Ports (naabu)", level=2)
    if naabu_results:
        document.add_paragraph(f"{len(naabu_results)} open port(s) found across the hosts above.")
        _add_table(document, ["Host", "Port"], naabu_results[:cap])
        if len(naabu_results) > cap:
            document.add_paragraph(f"... and {len(naabu_results) - cap} more (see naabu.log)")
    else:
        document.add_paragraph("No open ports were found on the scanned port list.")

    document.add_heading("C.4 Service and Vulnerability Scan (nmap + vulners)", level=2)
    if nmap_hosts:
        timed_out = [h for h in nmap_hosts if h["timeout"]]
        with_services = [h for h in nmap_hosts if h["ports"]]
        all_cves = sorted({c for h in nmap_hosts for c in h["cves"]})
        document.add_paragraph(
            f"{len(nmap_hosts)} host(s) were targeted: {len(with_services)} returned service "
            f"information, {len(timed_out)} were skipped because they exceeded the per-host "
            f"timeout and were therefore NOT assessed, and {len(all_cves)} distinct CVE "
            f"identifier(s) were matched by the vulners script."
        )
        rows = []
        for h in nmap_hosts[:cap]:
            if h["timeout"]:
                result = "Skipped (host timeout)"
            elif h["ports"]:
                result = "Scanned"
            else:
                result = "No open ports in scanned set"
            services = "; ".join(
                f"{p} {svc}{' ' + ver if ver else ''}" for p, svc, ver in h["ports"]) or "-"
            rows.append([h["host"], h["ip"], result, services, ", ".join(h["cves"][:10]) or "-"])
        _add_table(document, ["Host", "IP", "Result", "Services", "CVEs"], rows)
    elif nmap_raw_lines:
        document.add_paragraph("nmap ran but produced no host results.")
    else:
        document.add_paragraph("No nmap results were recorded for this target.")

    document.add_heading("C.5 Directory Enumeration (dirsearch)", level=2)
    ran = any(l.startswith("# Dirsearch started") for l in dirsearch_raw)
    if dirsearch_results:
        document.add_paragraph(f"{len(dirsearch_results)} interesting path(s) were found.")
        _add_table(document, ["Status", "Path"], dirsearch_results[:100])
    elif ran:
        document.add_paragraph(
            "dirsearch ran and found no paths with an interesting status code "
            "(200, 301, 302, 401, 403).")
    else:
        document.add_paragraph(
            "dirsearch.log contains no run record. The tool may not have executed for "
            "this scan, so absence of results here should not be read as 'no paths exist'.")

    document.add_heading("C.6 XSS Testing (dalfox)", level=2)
    record = [l.lstrip("# ").strip() for l in dalfox_raw if l.lstrip().startswith("#")]
    for l in record:
        document.add_paragraph(l)
    if dalfox_findings:
        document.add_paragraph(f"{len(dalfox_findings)} potential XSS finding line(s):")
        for l in dalfox_findings[:50]:
            document.add_paragraph(l)
    elif not record:
        document.add_paragraph(
            "dalfox.log contains no run record. The XSS stage may not have executed "
            "for this scan.")
    else:
        document.add_paragraph("No XSS findings were recorded.")

    document.save(out_path)


def build_templated_report(ctx_paths, args, prefix_name):
    if not HAVE_DOCXTPL:
        raise RuntimeError("docxtpl not installed")
    if not os.path.exists(TEMPLATE_PATH):
        raise RuntimeError(f"report template not found at {TEMPLATE_PATH}")

    f = ctx_paths

    # uniqdomains.log can hold the same host several times, some lines carrying
    # httpx status annotations ("host [200] [HSTS]"). Reduce to unique bare
    # hostnames so the counts and the appendix inventory are accurate.
    uniqdomains = []
    for line in read_lines(f("uniqdomains.log")):
        tokens = strip_ansi(line).split()
        if tokens and tokens[0] not in uniqdomains:
            uniqdomains.append(tokens[0])
    naabu_results = parse_naabu(read_lines(f("naabu.log")))
    nuclei_findings_raw = parse_nuclei_jsonl(read_lines(f("nucleiAlerts.log")))
    dirsearch_raw = read_lines(f("dirsearch.log"))
    dirsearch_results = parse_dirsearch(dirsearch_raw)
    # '#'-prefixed lines are phase-xss.sh's attempt record, not findings
    dalfox_raw = read_lines(f("dalfox.log"))
    dalfox_lines = [l for l in dalfox_raw if not l.lstrip().startswith("#")]
    unique_ips = [strip_ansi(l).strip() for l in read_lines(f("unique-ips.log")) if l.strip()]
    nmap_raw_lines = read_lines(f("nmapvulners.log"))
    nmap_hosts = parse_nmap(nmap_raw_lines)

    scope_file = os.path.join(args.target_dir, "scope.txt")
    scope_patterns = read_lines(scope_file)

    sev_counts = {sev: 0 for sev in SEVERITY_ORDER}
    for finding in nuclei_findings_raw:
        sev_counts[finding["severity"]] += 1
    total_findings = sum(sev_counts.values())

    findings_ctx = build_finding_context(nuclei_findings_raw)
    top_findings = [finding["name"] for finding in nuclei_findings_raw if finding["severity"] in ("critical", "high")]

    print(f" █▄▄▪ [report] {total_findings} nuclei findings, {len(naabu_results)} open ports, "
          f"{len(uniqdomains)} subdomains, {len(dirsearch_results)} dirsearch hits, "
          f"{len(dalfox_lines)} XSS candidates")

    print(f" █▄▄▪ [report] Generating executive summary (LLM: {args.llm_url})")
    exec_summary = generate_executive_summary(args.llm_url, args.llm_model, prefix_name, sev_counts, top_findings)

    if total_findings == 0:
        impact_statement = (
            "No nuclei findings were recorded for this target during the assessed window. "
            "This reflects what the automated toolchain detected and matched against its "
            "current template set — it is not a guarantee that no vulnerabilities exist."
        )
    else:
        crit_high = sev_counts.get("critical", 0) + sev_counts.get("high", 0)
        if crit_high > 0:
            impact_statement = (
                f"The {crit_high} critical/high severity finding(s) identified give potential "
                f"attackers the opportunity for unauthenticated access, data exposure, or service "
                f"disruption depending on the specific finding — see the Assessment Findings "
                f"section below for per-finding detail. In order to ensure data confidentiality, "
                f"integrity, and availability, remediations should be implemented as described in "
                f"the findings below, prioritized by severity."
            )
        else:
            impact_statement = (
                "No critical or high severity findings were identified. Medium/low/informational "
                "findings should still be reviewed and remediated where feasible, per the severity "
                "table above and the detail in the Assessment Findings section below."
            )

    surface = []
    if uniqdomains:
        surface.append(f"{len(uniqdomains)} hostname(s)")
    if unique_ips:
        surface.append(f"{len(unique_ips)} IP address(es)")
    if naabu_results:
        surface.append(f"{len(naabu_results)} open port(s)")
    if surface:
        impact_statement += (
            f" The assessed attack surface comprised {', '.join(surface)}; the full "
            f"inventory and per-tool results are listed in Appendix C."
        )

    methodology_text = (
        f"{args.team_name}'s testing methodology is an automated pipeline run by Kuromaku, split "
        f"into five sequential phases: Recon (subfinder/amass/gau/katana subdomain and URL "
        f"discovery), Network (dnsx DNS resolution and naabu port scanning), Webscan (nuclei "
        f"template matching across all 14 top-level categories, nmap+vulners service detection, "
        f"and dirsearch directory brute-forcing), XSS (dalfox testing of reflected parameters "
        f"found during recon), and Report (this document). IP/CIDR-only targets run a reduced "
        f"2-phase pipeline (Network then Webscan) since there are no hostnames for Recon/XSS to "
        f"operate on. Every phase is automated template/signature matching, not manual exploitation "
        f"or manual verification — findings should be triaged and confirmed before remediation or "
        f"disclosure."
    )

    tools_used = [
        {"name": "subfinder", "description": "Passive subdomain enumeration."},
        {"name": "amass", "description": "Passive subdomain enumeration (additional sources), bounded by its own timeout."},
        {"name": "gau / katana", "description": "Historical and JS-aware crawled URL discovery, feeding both domain enumeration and XSS candidate parameters."},
        {"name": "dnsx", "description": "DNS resolution of discovered hostnames to IPs."},
        {"name": "naabu", "description": "Port scanning (common web/service ports)."},
        {"name": "nuclei", "description": "Template-based vulnerability/exposure scanning across dns, iot, cves, technologies, exposures, fuzzing, miscellaneous, misconfiguration, default-logins, network, headless, takeovers, exposed-panels, and vulnerabilities categories."},
        {"name": "nmap + vulners", "description": "Service/version detection and known-CVE matching on ports naabu found open."},
        {"name": "dirsearch", "description": "Directory/file brute-forcing."},
        {"name": "dalfox", "description": "Reflected-XSS testing against parameters found with a reflection marker."},
    ]

    methodology_image_path = os.path.join(args.target_dir, "_methodology_diagram.png")
    diagram_ok = generate_methodology_diagram(methodology_image_path)

    doc = DocxTemplate(TEMPLATE_PATH)
    context = {
        "team_name": args.team_name,
        "client_name": args.client_name,
        "year": str(datetime.now().year),
        "test_date": datetime.now().strftime("%B %d, %Y"),
        "num_months": "3",
        "exec_summary": exec_summary,
        "impact_statement": impact_statement,
        "sev_counts": sev_counts,
        "scope_patterns": scope_patterns,
        "methodology_text": methodology_text,
        "findings": findings_ctx,
        "tools_used": tools_used,
        "methodology_image": (
            InlineImage(doc, methodology_image_path, width=Mm(180)) if diagram_ok else ""
        ),
    }
    # autoescape: values such as the "<CLIENT NAME>" default, nuclei matched
    # URLs and finding text contain <, > and &. Unescaped, they were written
    # into word/document.xml as raw tags, nesting paragraphs inside
    # paragraphs, which Word rejects as a corrupt file.
    doc.render(context, autoescape=True)

    out_path = f("report.docx")
    doc.save(out_path)

    # Never let post-processing take the report down — the base report is
    # already complete and saved.
    try:
        fixed = dedupe_paragraph_ids(out_path)
        if fixed:
            print(f" █▄▄▪ [report] Repaired {fixed} Word-compatibility issue(s) (unnamed styles / repeated paragraph ids)")
    except Exception as e:
        print(f" █▄▄▪ [report] paragraph-id cleanup skipped ({e})", file=sys.stderr)
    try:
        append_assessment_data(out_path, uniqdomains, unique_ips, naabu_results,
                               nmap_hosts, nmap_raw_lines, dirsearch_raw,
                               dirsearch_results, dalfox_raw, dalfox_lines)
    except Exception as e:
        print(f" █▄▄▪ [report] Appendix C not added ({e}); base report is intact", file=sys.stderr)
    return out_path


# ─── Legacy flat-summary report (fallback) ──────────────────────────────────
def add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def add_severity_table(doc, sev_counts):
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text = "Severity"
    hdr[1].text = "Count"
    for sev in SEVERITY_ORDER:
        count = sev_counts.get(sev, 0)
        if count == 0:
            continue
        row = table.add_row().cells
        run = row[0].paragraphs[0].add_run(sev.upper())
        run.bold = True
        run.font.color.rgb = SEVERITY_COLORS[sev]
        row[1].text = str(count)


def add_cover_table(doc, target, scan_date):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    rows_data = [
        ("Target", target),
        ("Assessment Type", "Automated Recon + Vulnerability Scan (Kuromaku)"),
        ("Scope", "Subdomain enumeration, port scan, nuclei templates, dirsearch, XSS probing"),
        ("Date", scan_date),
        ("Tooling", "subfinder, amass, httpx, gau, katana, dnsx, naabu, nuclei, nmap, dirsearch, dalfox"),
        ("Classification", "Internal Use Only"),
    ]
    for label, value in rows_data:
        row = table.add_row().cells
        row[0].text = label
        row[1].text = value
        row[0].paragraphs[0].runs[0].bold = True


def add_finding_section(doc, sev, findings, level=2):
    if not findings:
        return
    add_heading(doc, f"{sev.upper()} Findings ({len(findings)})", level=level)
    for f in findings:
        p = doc.add_paragraph()
        run = p.add_run(f"[{sev.upper()}] {f['name']}")
        run.bold = True
        run.font.color.rgb = SEVERITY_COLORS[sev]

        detail = doc.add_paragraph()
        detail.add_run("Target: ").bold = True
        detail.add_run(f.get('matched_at', ''))
        if f.get('cvss_score') is not None:
            detail.add_run("\nCVSS: ").bold = True
            detail.add_run(str(f['cvss_score']))
        detail.add_run("\nRecommendation: ").bold = True
        detail.add_run(
            f.get('remediation') or (
                "Review this finding against the template documentation at "
                "https://github.com/projectdiscovery/nuclei-templates and remediate "
                "per vendor guidance for the affected technology."
            )
        )


def add_finding_register(doc, all_findings_flat):
    add_heading(doc, "Consolidated Finding Register", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, h in enumerate(["#", "Finding", "Severity", "Target"]):
        hdr[i].text = h
        hdr[i].paragraphs[0].runs[0].bold = True

    for i, f in enumerate(all_findings_flat, 1):
        row = table.add_row().cells
        row[0].text = str(i)
        row[1].text = f["name"]
        sev_run = row[2].paragraphs[0].add_run(f["severity"].upper())
        sev_run.font.color.rgb = SEVERITY_COLORS.get(f["severity"], RGBColor(0, 0, 0))
        sev_run.bold = True
        row[3].text = f.get("matched_at", "")


def add_infra_overview(doc, naabu_results, resolved_lines):
    add_heading(doc, "Infrastructure Overview", level=1)
    if resolved_lines:
        add_heading(doc, "DNS Resolution", level=2)
        doc.add_paragraph(f"{len(resolved_lines)} DNS records resolved during enumeration.")
    if naabu_results:
        add_heading(doc, "Open Ports / Services", level=2)
        table = doc.add_table(rows=1, cols=2)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text = "Host"
        hdr[1].text = "Port"
        for host, port in naabu_results[:100]:
            row = table.add_row().cells
            row[0].text = host
            row[1].text = port
        if len(naabu_results) > 100:
            doc.add_paragraph(f"... and {len(naabu_results) - 100} more (see naabu.log)")


def add_appendix(doc, uniqdomains, dirsearch_results, dalfox_lines):
    add_heading(doc, "Appendix", level=1)
    if uniqdomains:
        add_heading(doc, "Subdomain Inventory", level=2)
        for d in uniqdomains[:200]:
            doc.add_paragraph(d, style="List Bullet")
        if len(uniqdomains) > 200:
            doc.add_paragraph(f"... and {len(uniqdomains) - 200} more (see uniqdomains.log)")
    if dirsearch_results:
        add_heading(doc, "Directory Enumeration — Interesting Paths", level=2)
        table = doc.add_table(rows=1, cols=2)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text = "Status"
        hdr[1].text = "Path"
        for status, path in dirsearch_results[:100]:
            row = table.add_row().cells
            row[0].text = status
            row[1].text = path
    if dalfox_lines:
        add_heading(doc, "XSS Candidate Findings (dalfox)", level=2)
        for line in dalfox_lines[:50]:
            doc.add_paragraph(line, style="List Bullet")


def build_legacy_report(ctx_paths, args, prefix_name):
    if not HAVE_DOCX:
        raise RuntimeError("python-docx not installed")

    f = ctx_paths
    uniqdomains = read_lines(f("uniqdomains.log"))
    naabu_results = parse_naabu(read_lines(f("naabu.log")))
    resolved_lines = read_lines(f("resolved.log"))
    nuclei_findings_raw = parse_nuclei_jsonl(read_lines(f("nucleiAlerts.log")))
    dirsearch_results = parse_dirsearch(read_lines(f("dirsearch.log")))
    # '#'-prefixed lines are phase-xss.sh's attempt record, not findings
    dalfox_lines = [l for l in read_lines(f("dalfox.log")) if not l.lstrip().startswith("#")]

    nuclei_by_sev = defaultdict(list)
    for finding in nuclei_findings_raw:
        nuclei_by_sev[finding["severity"]].append(finding)
    sev_counts = {sev: len(nuclei_by_sev.get(sev, [])) for sev in SEVERITY_ORDER}
    total_findings = sum(sev_counts.values())

    all_findings_flat = []
    for sev in SEVERITY_ORDER:
        for fnd in nuclei_by_sev.get(sev, []):
            all_findings_flat.append(fnd)

    top_findings = [fnd["name"] for sev in ("critical", "high") for fnd in nuclei_by_sev.get(sev, [])]

    print(f" █▄▄▪ [report] {total_findings} nuclei findings, {len(naabu_results)} open ports, "
          f"{len(uniqdomains)} subdomains, {len(dirsearch_results)} dirsearch hits, "
          f"{len(dalfox_lines)} XSS candidates")
    print(f" █▄▄▪ [report] Generating executive summary (LLM: {args.llm_url})")
    exec_summary = generate_executive_summary(args.llm_url, args.llm_model, prefix_name, sev_counts, top_findings)

    doc = Document()
    doc.add_heading("AUTOMATED RECONNAISSANCE & VULNERABILITY REPORT", level=0)
    doc.add_paragraph(prefix_name).alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_heading(doc, "Assessment Details", level=1)
    add_cover_table(doc, prefix_name, datetime.now().strftime("%B %d, %Y"))

    add_heading(doc, "Risk Finding Summary", level=1)
    add_severity_table(doc, sev_counts)

    add_heading(doc, "1. Executive Summary", level=1)
    doc.add_paragraph(exec_summary)

    add_infra_overview(doc, naabu_results, resolved_lines)

    add_heading(doc, "Findings — Ranked by Risk", level=1)
    doc.add_paragraph(
        "Findings below are derived from automated nuclei template matches, "
        "ranked by severity as classified by the matched template."
    )
    if total_findings == 0:
        doc.add_paragraph("No nuclei findings were recorded for this target.")
    for sev in SEVERITY_ORDER:
        add_finding_section(doc, sev, nuclei_by_sev.get(sev, []))

    if all_findings_flat:
        add_finding_register(doc, all_findings_flat)

    add_appendix(doc, uniqdomains, dirsearch_results, dalfox_lines)

    add_heading(doc, "Disclaimer", level=1)
    doc.add_paragraph(
        "This report was generated automatically by the Kuromaku scanning "
        "pipeline. Findings are based on automated tooling and template "
        "matching; manual verification is recommended before remediation "
        "or disclosure. This document may contain sensitive information "
        "about the target's attack surface and should be handled accordingly."
    )

    out_path = f("report.docx")
    doc.save(out_path)
    return out_path


# ─── Main ────────────────────────────────────────────────────────────────────
def main():
    args = parse_args()
    tdir = args.target_dir

    if args.domain_list:
        prefix_name = os.path.basename(args.domain_list)
    else:
        prefix_name = args.domain

    if not prefix_name:
        print("Error: -d or -dl required", file=sys.stderr)
        sys.exit(1)

    target_prefix = os.path.join(tdir, prefix_name)

    def f(suffix):
        return f"{target_prefix}-{suffix}"

    print(f" █▄▄▪ [report] Reading scan outputs for {prefix_name}")

    out_path = None
    if not args.legacy:
        try:
            out_path = build_templated_report(f, args, prefix_name)
            print(f" █▄▄▪ [report] Templated report written: {out_path}")
        except Exception as e:
            print(f" █▄▄▪ [report] Templated report generation failed ({e}) — falling back to legacy format", file=sys.stderr)
            out_path = None

    if out_path is None:
        out_path = build_legacy_report(f, args, prefix_name)
        print(f" █▄▄▪ [report] Legacy report written: {out_path}")


if __name__ == "__main__":
    main()
